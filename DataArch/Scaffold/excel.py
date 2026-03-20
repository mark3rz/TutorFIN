"""
parsers/excel.py — Excel / CSV ingestion and extraction.
Uses openpyxl for .xlsx, csv module for .csv.
"""

import csv
import openpyxl
from pathlib import Path
from pipeline.schema import ParsedDocument, DocumentMetadata, DocumentType, ExtractionStatus, BusinessEntity
from pipeline import llm


def _xlsx_to_text(file_path: Path) -> tuple[str, int]:
    """Returns (text_repr, sheet_count)."""
    wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
    sheets = wb.sheetnames
    lines = []
    for sheet_name in sheets:
        ws = wb[sheet_name]
        lines.append(f"=== Sheet: {sheet_name} ===")
        for row in ws.iter_rows(values_only=True):
            row_text = "\t".join(str(c) if c is not None else "" for c in row)
            if row_text.strip():
                lines.append(row_text)
    wb.close()
    return "\n".join(lines), len(sheets)


def _csv_to_text(file_path: Path) -> str:
    with open(file_path, newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        return "\n".join("\t".join(row) for row in reader)


def parse(file_path: Path) -> ParsedDocument:
    suffix = file_path.suffix.lower()
    doc_type = DocumentType.EXCEL

    metadata = DocumentMetadata(
        filename=file_path.name,
        file_type=doc_type,
        file_size_bytes=file_path.stat().st_size,
        source_path=str(file_path),
    )

    errors = []
    raw_text = ""

    try:
        if suffix in (".xlsx", ".xlsm"):
            raw_text, sheet_count = _xlsx_to_text(file_path)
            metadata.sheet_count = sheet_count
        elif suffix == ".csv":
            raw_text = _csv_to_text(file_path)
        else:
            errors.append(f"Unsupported Excel format: {suffix}")
            return ParsedDocument(metadata=metadata, status=ExtractionStatus.FAILED, errors=errors)
    except Exception as e:
        errors.append(f"Excel parse error: {e}")
        return ParsedDocument(metadata=metadata, status=ExtractionStatus.FAILED, errors=errors)

    try:
        extracted = llm.extract(raw_text)
        entities = [BusinessEntity(**e) for e in extracted.get("entities", [])]
        return ParsedDocument(
            metadata=metadata,
            status=ExtractionStatus.SUCCESS,
            raw_text=raw_text,
            entities=entities,
            summary=extracted.get("summary", ""),
            data_fields=extracted.get("data_fields", {}),
            errors=errors,
        )
    except Exception as e:
        errors.append(f"LLM extraction error: {e}")
        return ParsedDocument(metadata=metadata, status=ExtractionStatus.PARTIAL, raw_text=raw_text, errors=errors)
