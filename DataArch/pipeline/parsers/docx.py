"""
parsers/docx.py — Word document ingestion and extraction.
Uses python-docx for text extraction, Claude for structured data.
"""

import docx
from pathlib import Path
from pipeline.schema import ParsedDocument, DocumentMetadata, DocumentType, ExtractionStatus, BusinessEntity
from pipeline import llm


def _docx_to_text(file_path: Path) -> tuple[str, int]:
    """Returns (full_text, paragraph_count)."""
    doc = docx.Document(str(file_path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also extract table content
    table_lines = []
    for table in doc.tables:
        for row in table.rows:
            row_text = "\t".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                table_lines.append(row_text)

    all_text = "\n".join(paragraphs)
    if table_lines:
        all_text += "\n\n=== Tables ===\n" + "\n".join(table_lines)

    return all_text, len(paragraphs)


def parse(file_path: Path) -> ParsedDocument:
    metadata = DocumentMetadata(
        filename=file_path.name,
        file_type=DocumentType.DOCX,
        file_size_bytes=file_path.stat().st_size,
        source_path=str(file_path),
    )

    errors = []
    raw_text = ""

    try:
        raw_text, para_count = _docx_to_text(file_path)
        metadata.page_count = None  # python-docx doesn't expose page count easily
    except Exception as e:
        errors.append(f"DOCX parse error: {e}")
        return ParsedDocument(metadata=metadata, status=ExtractionStatus.FAILED, errors=errors)

    if not raw_text.strip():
        errors.append("No extractable text found in document.")
        return ParsedDocument(metadata=metadata, status=ExtractionStatus.FAILED, raw_text=raw_text, errors=errors)

    try:
        extracted = llm.extract(raw_text)
        entities = [BusinessEntity(**e) for e in extracted.get("entities", [])]
        return ParsedDocument(
            metadata=metadata,
            status=ExtractionStatus.SUCCESS,
            raw_text=raw_text,
            entities=entities,
            summary=extracted.get("summary", ""),
            data_fields=extracted.get("data_fields", {}),
            errors=errors,
        )
    except Exception as e:
        errors.append(f"LLM extraction error: {e}")
        return ParsedDocument(metadata=metadata, status=ExtractionStatus.PARTIAL, raw_text=raw_text, errors=errors)
